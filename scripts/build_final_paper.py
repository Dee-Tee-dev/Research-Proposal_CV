from __future__ import annotations

import csv
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper/final_short_paper.md"
OUTPUT_DIR = ROOT / "output"
DOCX_PATH = OUTPUT_DIR / "Income_Gaps_VLM_Short_Paper.docx"
PDF_PATH = OUTPUT_DIR / "pdf/Income_Gaps_VLM_Short_Paper.pdf"
ASSET_DIR = ROOT / "paper/assets/combined"
CHART_PATH = ASSET_DIR / "scores_by_income.png"
QUALITATIVE_PATH = ASSET_DIR / "qualitative_examples.png"


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def set_cell_margins(cell, top=55, start=65, bottom=55, end=65) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color="D9D9D9", size="4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def clean_inline(text: str) -> str:
    return text.replace("**", "").replace("`", "")


def add_body_paragraph(doc: Document, text: str, italic=False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(2.5)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = False
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(clean_inline(text))
    run.italic = italic


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    data = [[clean_inline(cell.strip()) for cell in line.strip().strip("|").split("|")] for line in lines]
    data = [data[0]] + data[2:]
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if len(data[0]) == 6:
        widths = [Inches(2.25)] + [Inches(0.72)] * 5
    elif len(data[0]) == 4:
        widths = [Inches(1.05), Inches(1.45), Inches(1.85), Inches(2.85)]
    else:
        widths = [Inches(7.2 / len(data[0]))] * len(data[0])
    for row_idx, row in enumerate(data):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.width = widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            if row_idx == 0:
                set_cell_shading(cell, "17365D")
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "EAF1F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.name = "Arial"
            run.font.size = Pt(6.8 if len(data[0]) == 6 else 7.1)
            if row_idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_docx() -> None:
    if not CHART_PATH.exists():
        raise FileNotFoundError(f"Create the combined chart first: {CHART_PATH}")
    if not QUALITATIVE_PATH.exists():
        raise FileNotFoundError(f"Create the qualitative figure first: {QUALITATIVE_PATH}")
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8.75)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.0
    for name, size in (("Title", 15), ("Heading 1", 11), ("Heading 2", 9.2), ("Heading 3", 8.6)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(4 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.keep_with_next = True
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    idx = 0
    inserted_chart = False
    while idx < len(lines):
        line = lines[idx].strip()
        if not line:
            idx += 1
            continue
        if line.startswith("|") and idx + 1 < len(lines) and "---" in lines[idx + 1]:
            table_lines = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                table_lines.append(lines[idx])
                idx += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(clean_inline(line[2:]))
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(15)
        elif line.startswith("## Appendix A"):
            doc.add_page_break()
            doc.add_heading(clean_inline(line[3:]), level=1)
        elif line.startswith("## "):
            doc.add_heading(clean_inline(line[3:]), level=1)
        elif line.startswith("### "):
            doc.add_heading(clean_inline(line[4:]), level=2)
        elif line.startswith("**Diya"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(clean_inline(line))
            run.bold = True
            run.font.size = Pt(9.5)
        elif line.startswith("!["):
            picture = doc.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.space_after = Pt(1)
            picture.paragraph_format.keep_with_next = True
            picture.add_run().add_picture(str(QUALITATIVE_PATH), width=Inches(6.9))
        elif line.startswith("*") and line.endswith("*"):
            add_body_paragraph(doc, line.strip("*"), italic=True)
            if line.startswith("*Table 1") and not inserted_chart:
                picture = doc.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.paragraph_format.keep_with_next = True
                picture.add_run().add_picture(str(CHART_PATH), width=Inches(6.9))
                caption = doc.add_paragraph()
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.paragraph_format.space_after = Pt(2)
                run = caption.add_run(
                    "Figure 1. Scores by income quartile. Each point represents 42 images; task-specific metrics are not interchangeable."
                )
                run.italic = True
                run.font.size = Pt(7.2)
                inserted_chart = True
        elif line.startswith(tuple(f"{n}. " for n in range(1, 10))):
            paragraph = doc.add_paragraph(style="Normal")
            paragraph.paragraph_format.left_indent = Inches(0.12)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12)
            paragraph.paragraph_format.space_after = Pt(1)
            run = paragraph.add_run(clean_inline(line))
            run.font.size = Pt(7.5)
        elif line == "Computer Vision Project Short Paper":
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(5)
            run = paragraph.add_run(line)
            run.font.size = Pt(8)
        else:
            add_body_paragraph(doc, line)
        idx += 1

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    PDF_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def convert_pdf() -> None:
    subprocess.run(
        [
            "/Users/divyatiwari/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(PDF_PATH.parent),
            str(DOCX_PATH),
        ],
        check=True,
    )


if __name__ == "__main__":
    build_docx()
    convert_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
